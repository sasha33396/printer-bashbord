"""Create-only, atomic import of the two digital document registers from Excel."""
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet._reader import WorkSheetParser
from pydantic import ValidationError
from sqlalchemy import text

from digital_document_schemas import SignatureCreate, PowerOfAttorneyCreate
from models import ElectronicSignature, MachineReadablePowerOfAttorney


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
LOCAL_TIMEZONE = timezone(timedelta(hours=4))
FIELDS = {
    'ecp': [
        ('company_name', 'Название компании'), ('inn', 'ИНН'), ('certificate_type', 'Тип'),
        ('signature_kind', 'Вид подписи'), ('full_name', 'ФИО'), ('position', 'Должность'),
        ('snils', 'СНИЛС'), ('email', 'Email'), ('application', 'Применение'),
        ('valid_from', 'Действителен с'), ('valid_to', 'Действителен по'),
        ('ep_state', 'Состояние ЭП'), ('revoked_at', 'Дата и время отзыва'),
        ('fingerprint', 'Отпечаток'), ('serial_number', 'Серийный номер'),
    ],
    'mchd': [
        ('power_number', 'Номер доверенности'), ('valid_from', 'Дата начала действия'),
        ('valid_to', 'Дата окончания действия'), ('grantor_inn', 'ИНН доверителя'),
        ('grantor_name', 'Наименование доверителя'), ('grantor_person_full_name', 'ФИО лица доверителя'),
        ('grantor_person_inn', 'ИНН лица доверителя'), ('grantor_person_snils', 'СНИЛС лица доверителя'),
        ('representative_full_name', 'ФИО лица представителя'), ('representative_inn', 'ИНН лица представителя'),
        ('representative_snils', 'СНИЛС лица представителя'), ('permissions', 'Разрешения'),
        ('fns_identifier', 'Идентификатор МЧД ФНС'), ('edo_identifier', 'Идентификатор МЧД ЭДО'),
        ('edo_status', 'Статус МЧД ЭДО'),
    ],
}
REGISTERS = {
    'ecp': ('ЭЦП', ElectronicSignature, SignatureCreate),
    'mchd': ('МЧД', MachineReadablePowerOfAttorney, PowerOfAttorneyCreate),
}
IDENTIFIERS = {'inn', 'snils', 'serial_number', 'fingerprint', 'power_number', 'fns_identifier',
               'edo_identifier', 'grantor_inn', 'grantor_person_inn', 'grantor_person_snils',
               'representative_inn', 'representative_snils'}


def header(value):
    normalized = ' '.join(str(value or '').replace('\ufeff', '').split()).strip().casefold()
    return 'действителен с' if normalized == 'действителен c' else normalized


def nonempty_rows(sheet):
    """Read physical rows, ignoring unreliable dimensions and empty formatting.

    Use the streaming cell parser from the pinned openpyxl version. iter_rows()
    trusts dimension metadata and generates every missing row up to sparse cells;
    a styled cell at XFD1048576 must not create a million rows to validate.
    """
    with sheet._get_source() as source:
        parser = WorkSheetParser(source, sheet._shared_strings,
                                 data_only=False, epoch=sheet.parent.epoch,
                                 date_formats=sheet.parent._date_formats)
        for row_number, cells in parser.parse():
            populated = {
                cell['column']: ReadOnlyCell(sheet, **cell) for cell in cells
                if cell['value'] is not None
                and not (isinstance(cell['value'], str) and not cell['value'].strip())
            }
            if populated:
                yield row_number, populated


def excel_date(value, epoch):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        value = value.strip()
        try:
            return datetime.fromisoformat(value.replace('Z', '+00:00'))
        except ValueError:
            for pattern in ('%d.%m.%Y', '%d.%m.%Y %H:%M', '%d.%m.%Y %H:%M:%S'):
                try:
                    return datetime.strptime(value, pattern)
                except ValueError:
                    pass
    raise ValueError('Укажите дату Excel, ДД.ММ.ГГГГ или дату в формате ISO')


def cell_value(cell, key, epoch):
    value = cell.value
    if cell.data_type == 'f':
        raise ValueError('Формула: замените её результатом через «Вставить значения» в Excel')
    if cell.data_type == 'e':
        raise ValueError('Ячейка содержит ошибку Excel')
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if key in {'valid_from', 'valid_to', 'revoked_at'}:
        if isinstance(value, str) and value.strip() in {'-', '—', '–'}:
            return None
        parsed = excel_date(value, epoch)
        if key != 'revoked_at':
            return parsed.date()
        return (parsed if parsed.tzinfo else parsed.replace(tzinfo=LOCAL_TIMEZONE)).astimezone(timezone.utc)
    if isinstance(value, bool):
        raise ValueError('Ожидается текст, а не логическое значение')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or int(value) != value:
            raise ValueError('Укажите значение текстом, без дробной части')
        result = str(int(value))
        if key in IDENTIFIERS:
            if len(result.lstrip('-')) > 15:
                raise ValueError('Длинный номер записан числом: Excel мог потерять цифры. Укажите исходный номер текстом')
            # Recover only explicitly displayed leading zeroes, never guess identifier lengths.
            if re.fullmatch(r'0+', cell.number_format):
                result = result.zfill(len(cell.number_format))
        return result
    if not isinstance(value, str):
        raise ValueError('Ожидается текстовое значение')
    return value.strip()


def parse_workbook(contents):
    if len(contents) > MAX_FILE_BYTES:
        raise ValueError('Размер файла не должен превышать 10 МБ')
    try:
        with ZipFile(BytesIO(contents)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Слишком большой объём данных внутри XLSX')
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException) as exc:
        raise ValueError('Не удалось прочитать XLSX. Сохраните файл в формате .xlsx') from exc
    rows, warnings = [], []
    try:
        for kind, (sheet_name, _, schema) in REGISTERS.items():
            matches = [sheet for sheet in workbook if header(sheet.title) == header(sheet_name)]
            if len(matches) != 1:
                raise ValueError(f'В файле должен быть один лист «{sheet_name}»')
            sheet = matches[0]
            labels = dict(FIELDS[kind])
            names = {header(label): key for key, label in FIELDS[kind]}
            required = {key for key, field in schema.model_fields.items() if field.is_required()}
            columns, header_row = {}, None
            # Allow a title above the table; identify its header by all required columns.
            header_rows = nonempty_rows(sheet)
            try:
                for row_number, cells in header_rows:
                    if row_number > 20:
                        break
                    recognized = [(column, names[header(cell.value)]) for column, cell in cells.items() if header(cell.value) in names]
                    if required.issubset({key for _, key in recognized}):
                        if len(recognized) != len({key for _, key in recognized}):
                            raise ValueError(f'Лист «{sheet_name}»: заголовки столбцов повторяются')
                        columns, header_row = dict(recognized), row_number
                        break
            finally:
                header_rows.close()
            if header_row is None:
                needed = ', '.join(labels[key] for key in labels if key in required)
                raise ValueError(f'Лист «{sheet_name}»: не найдены заголовки в первых 20 строках. Обязательные столбцы: {needed}')
            missing = [label for key, label in FIELDS[kind] if key not in columns.values()]
            if missing:
                warnings.append(f'{sheet_name}: нет необязательных столбцов ({", ".join(missing)}), значения будут пустыми')
            count = 0
            for row_number, cells in nonempty_rows(sheet):
                if row_number <= header_row:
                    continue
                selected = [(columns[column], cell) for column, cell in cells.items() if column in columns]
                if not selected:
                    continue
                count += 1
                if count > MAX_ROWS:
                    raise ValueError(f'Лист «{sheet_name}»: максимум {MAX_ROWS} строк данных')
                entry = {'kind': kind, 'sheet': sheet_name, 'row': row_number, 'status': 'new',
                         'message': '', 'data': None, 'source_data': {}}
                values, errors = {}, []
                for key, cell in selected:
                    raw = cell.value
                    entry['source_data'][key] = raw.isoformat() if isinstance(raw, (date, datetime)) else str(raw)
                    try:
                        value = cell_value(cell, key, workbook.epoch)
                        values[key] = value
                        # Keep successfully read identifiers (including zero masks)
                        # even when a different field invalidates this row.
                        entry['source_data'][key] = value.isoformat() if isinstance(value, (date, datetime)) else value
                    except (ValueError, OverflowError) as exc:
                        errors.append(f'{labels[key]}: {exc}')
                if not errors:
                    try:
                        payload = schema.model_validate(values)
                        entry['data'] = payload.model_dump(mode='json')
                    except ValidationError as exc:
                        for error in exc.errors(include_url=False, include_input=False, include_context=False):
                            key = error['loc'][0] if error['loc'] else None
                            label = labels.get(key, 'Строка')
                            msg = error['msg']
                            if error['type'] in {'missing', 'string_type', 'date_type'}:
                                msg = 'Заполните обязательное поле' if key in required else 'Некорректное значение'
                            errors.append(f'{label}: {msg}')
                if errors:
                    entry.update(status='error', message='; '.join(errors))
                rows.append(entry)
        if not rows:
            raise ValueError('На листах «ЭЦП» и «МЧД» нет строк данных')
        return rows, warnings
    except (BadZipFile, ParseError, KeyError, OSError, InvalidFileException) as exc:
        raise ValueError('Не удалось прочитать таблицы XLSX. Пересохраните файл в Excel') from exc
    finally:
        workbook.close()


def normalized(value):
    return ' '.join(str(value or '').split()).casefold()


def identity_keys(kind, data):
    keys = []
    if kind == 'ecp':
        if data.get('fingerprint'):
            keys.append(('fingerprint', re.sub(r'[\s:]', '', data['fingerprint']).casefold()))
        if data.get('serial_number'):
            issuer = ('inn', normalized(data['inn'])) if data.get('inn') else ('company', normalized(data['company_name']))
            keys.append(('serial', issuer, normalized(data['serial_number'])))
    else:
        for key in ('fns_identifier', 'edo_identifier'):
            if data.get(key):
                keys.append((key, normalized(data[key])))
        # Number alone is not globally unique: include grantor and validity period.
        issuer = ('inn', normalized(data['grantor_inn'])) if data.get('grantor_inn') else ('company', normalized(data['grantor_name']))
        keys.append(('number', issuer, normalized(data['power_number']), data['valid_from'], data['valid_to']))
    return keys


def canonical(data):
    return tuple(sorted(data.items()))


def identity_summary(tokens, existing):
    """Bound conflict messages and avoid scanning a duplicate group per row."""
    first = canonical(existing[next(iter(tokens))][0])
    same_data = all(canonical(existing[token][0]) == first for token in tokens)
    return {'canonical': first if same_data else None,
            'db': sorted(token for token in tokens if token[0] == 'db')[:2],
            'file_rows': sorted(token[1] for token in tokens if token[0] == 'file')[:6]}


def import_documents(db, contents, apply=False):
    rows, warnings = parse_workbook(contents)
    # Serialize SQLite imports before checking identities; retrying the same file is safe.
    if apply and db.get_bind().dialect.name == 'sqlite':
        db.execute(text('BEGIN IMMEDIATE'))
    indexes = {}
    for kind, (_, model, schema) in REGISTERS.items():
        by_key, exact, existing = defaultdict(set), defaultdict(set), {}
        for record in db.query(model).all():
            data = {key: getattr(record, key) for key in schema.model_fields}
            if isinstance(data.get('revoked_at'), datetime) and data['revoked_at'].tzinfo is None:
                data['revoked_at'] = data['revoked_at'].replace(tzinfo=timezone.utc)
            data = schema.model_validate(data).model_dump(mode='json')
            token = ('db', record.id)
            existing[token] = (data, record)
            exact[canonical(data)].add(token)
            for key in identity_keys(kind, data):
                by_key[key].add(token)
        indexes[kind] = (by_key, exact, existing)
    # Index every valid file row first. Different rows sharing an identity must
    # both be withheld; partial import must not pick whichever appears first.
    for entry in rows:
        if entry['status'] == 'error':
            continue
        kind, data = entry['kind'], entry['data']
        by_key, exact, existing = indexes[kind]
        token = ('file', entry['row'])
        existing[token] = (data, entry)
        exact[canonical(data)].add(token)
        for key in identity_keys(kind, data):
            by_key[key].add(token)
    for kind, (by_key, exact, existing) in indexes.items():
        indexes[kind] = ({key: identity_summary(tokens, existing) for key, tokens in by_key.items()},
                         {key: identity_summary(tokens, existing) for key, tokens in exact.items()}, existing)
    pending = []
    for entry in rows:
        if entry['status'] == 'error':
            continue
        kind, data = entry['kind'], entry['data']
        by_key, exact, existing = indexes[kind]
        fingerprint = canonical(data)
        matches = [exact[fingerprint]]
        keys = identity_keys(kind, data)
        for key in keys:
            matches.append(by_key[key])
        db_candidates = list({token for match in matches for token in match['db']})
        if len(db_candidates) == 1:
            entry['existing_id'] = db_candidates[0][1]
        different = any(match['canonical'] != fingerprint for match in matches)
        if different or len(db_candidates) > 1:
            file_rows = sorted({number for match in matches for number in match['file_rows'] if number != entry['row']})
            message = 'Номер или идентификатор уже встречается с другими данными. Сверьте записи; автоматическая замена отключена'
            if file_rows:
                message += f'. Связанные строки этого листа: {", ".join(map(str, file_rows[:5]))}'
                if len(file_rows) > 5:
                    message += ', …'
            entry.update(status='conflict', message=message)
            continue
        if db_candidates:
            record = existing[db_candidates[0]][1]
            message = 'Уже есть в приложении'
            if record.is_archived:
                message += ' (в архиве; останется в архиве)'
            entry.update(status='unchanged', message=message)
            continue
        first_row = min(number for match in matches for number in match['file_rows'])
        if first_row != entry['row']:
            entry.update(status='unchanged', message=f'Повтор строки {first_row} в файле')
            continue
        pending.append(entry)
        if kind == 'ecp' and not keys:
            warnings.append(f'ЭЦП, строка {entry["row"]}: нет отпечатка и серийного номера; повтор определяется только по совпадению всех полей')
    can_apply = bool(pending)
    applied = apply and can_apply
    if applied:
        try:
            created = []
            for entry in pending:
                _, model, schema = REGISTERS[entry['kind']]
                record = model(**schema.model_validate(entry['data']).model_dump())
                db.add(record)
                created.append((entry, record))
            db.flush()
            identifiers = [record.id for _, record in created]
            db.commit()
            for (entry, _), identifier in zip(created, identifiers):
                entry.update(status='created', message='Добавлено в приложение', existing_id=identifier)
        except Exception:
            db.rollback()
            raise
    elif apply:
        db.rollback()
    totals = dict(Counter(entry['status'] for entry in rows))
    return {'mode': 'apply' if apply else 'preview', 'applied': applied, 'can_apply': can_apply,
            'totals': {key: totals.get(key, 0) for key in ('new', 'created', 'unchanged', 'error', 'conflict')},
            'rows': rows, 'warnings': warnings}
