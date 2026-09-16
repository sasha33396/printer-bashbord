from __future__ import annotations

import io
from datetime import date
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from database import get_db
from models import (
    Department, Device, DeviceStatus,
    RepairRecord, RepairType,
    ConsumableLog,
)
from routers.auth import get_current_user

router = APIRouter()

_auth = Depends(get_current_user)


# ---------------------------------------------------------------------------
# Response schemas
# ---------------------------------------------------------------------------

class DeviceCostRow(BaseModel):
    device_id: int
    inventory_number: str
    manufacturer: Optional[str] = None
    model: str
    device_type: Optional[str] = None
    department: Optional[str] = None
    branch: Optional[str] = None
    location: Optional[str] = None
    total_repair_cost: float
    repair_count: int
    total_consumable_cost: float
    consumable_count: int
    total_cost: float
    cost_per_page: Optional[float] = None


class RepairsByType(BaseModel):
    planned: int
    unplanned: int
    warranty: int


class SummaryResponse(BaseModel):
    total_devices_active: int
    total_cost_period: float
    top5_expensive: List[DeviceCostRow]
    repairs_by_type: RepairsByType


# ---------------------------------------------------------------------------
# Core aggregation helper
# ---------------------------------------------------------------------------

def _cost_rows(
    db: Session,
    date_from: Optional[date],
    date_to: Optional[date],
    branch_id: Optional[int],
    department_id: Optional[int],
) -> List[DeviceCostRow]:
    # --- Devices + optional filters ---
    dev_q = db.query(Device).options(
        joinedload(Device.department).joinedload(Department.branch)
    )
    if department_id is not None:
        dev_q = dev_q.filter(Device.department_id == department_id)
    if branch_id is not None:
        dept_ids = [
            row.id
            for row in db.query(Department.id)
            .filter(Department.branch_id == branch_id)
            .all()
        ]
        dev_q = dev_q.filter(Device.department_id.in_(dept_ids))

    devices = dev_q.all()
    if not devices:
        return []

    device_ids = [d.id for d in devices]

    # --- Repair aggregates per device ---
    r_q = (
        db.query(
            RepairRecord.device_id,
            func.coalesce(func.sum(RepairRecord.cost), 0.0).label("total_cost"),
            func.count(RepairRecord.id).label("cnt"),
            func.max(RepairRecord.page_counter).label("max_page"),
        )
        .filter(RepairRecord.device_id.in_(device_ids))
    )
    if date_from:
        r_q = r_q.filter(RepairRecord.date >= date_from)
    if date_to:
        r_q = r_q.filter(RepairRecord.date <= date_to)
    repair_map = {
        row.device_id: row
        for row in r_q.group_by(RepairRecord.device_id).all()
    }

    # --- Consumable aggregates per device ---
    c_q = (
        db.query(
            ConsumableLog.device_id,
            func.coalesce(
                func.sum(ConsumableLog.unit_cost * ConsumableLog.quantity), 0.0
            ).label("total_cost"),
            func.count(ConsumableLog.id).label("cnt"),
            func.max(ConsumableLog.page_counter).label("max_page"),
        )
        .filter(ConsumableLog.device_id.in_(device_ids))
    )
    if date_from:
        c_q = c_q.filter(ConsumableLog.date >= date_from)
    if date_to:
        c_q = c_q.filter(ConsumableLog.date <= date_to)
    cons_map = {
        row.device_id: row
        for row in c_q.group_by(ConsumableLog.device_id).all()
    }

    # --- Assemble rows ---
    rows: List[DeviceCostRow] = []
    for dev in devices:
        r = repair_map.get(dev.id)
        c = cons_map.get(dev.id)

        total_repair_cost = float(r.total_cost) if r else 0.0
        repair_count = int(r.cnt) if r else 0
        total_consumable_cost = float(c.total_cost) if c else 0.0
        consumable_count = int(c.cnt) if c else 0
        total_cost = total_repair_cost + total_consumable_cost

        # Берём максимальный счётчик страниц из обеих таблиц
        pages = [p for p in (
            r.max_page if r else None,
            c.max_page if c else None,
        ) if p is not None]
        max_page = max(pages) if pages else None
        cost_per_page = round(total_cost / max_page, 4) if max_page else None

        dept = dev.department
        rows.append(DeviceCostRow(
            device_id=dev.id,
            inventory_number=dev.inventory_number,
            manufacturer=dev.manufacturer,
            model=dev.model,
            device_type=dev.device_type.value,
            department=dept.name if dept else None,
            branch=dept.branch.name if dept and dept.branch else None,
            location=dev.location,
            total_repair_cost=round(total_repair_cost, 2),
            repair_count=repair_count,
            total_consumable_cost=round(total_consumable_cost, 2),
            consumable_count=consumable_count,
            total_cost=round(total_cost, 2),
            cost_per_page=cost_per_page,
        ))

    return rows


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

_TYPE_RU: dict[str, str] = {
    "printer": "Принтер",
    "mfc":     "МФУ",
    "plotter": "Плоттер",
    "scanner": "Сканер",
}

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="4472C4")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def _style_header_row(ws) -> None:
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN


@router.get("/costs", response_model=List[DeviceCostRow])
def costs(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    branch_id: Optional[int] = None,
    department_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _: dict = _auth,
):
    rows = _cost_rows(db, date_from, date_to, branch_id, department_id)
    return sorted(rows, key=lambda r: r.total_cost, reverse=True)


@router.get("/costs/export")
def export_costs_detailed(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    branch_id: Optional[int] = None,
    department_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _: dict = _auth,
):
    rows = _cost_rows(db, date_from, date_to, branch_id, department_id)
    rows = sorted(rows, key=lambda r: r.total_cost, reverse=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Затраты по оборудованию"

    ws.append([
        "Инв.№", "Производитель", "Модель", "Тип",
        "Филиал", "Отдел", "Кабинет",
        "Затраты ремонт", "Кол-во ремонтов",
        "Затраты расходники", "Итого", "Цена отпечатка",
    ])
    _style_header_row(ws)

    for r in rows:
        ws.append([
            r.inventory_number,
            r.manufacturer or "",
            r.model,
            _TYPE_RU.get(r.device_type or "", r.device_type or ""),
            r.branch or "",
            r.department or "",
            r.location or "",
            r.total_repair_cost,
            r.repair_count,
            r.total_consumable_cost,
            r.total_cost,
            r.cost_per_page if r.cost_per_page is not None else "",
        ])

    # Числовой формат для денежных колонок (H, J, K)
    money_fmt = '#,##0.00 ₽'
    for col_letter in ("H", "J", "K"):
        for cell in ws[col_letter][1:]:
            cell.number_format = money_fmt

    col_widths = [14, 18, 20, 10, 20, 20, 10, 18, 16, 20, 14, 16]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_from_str = str(date_from) if date_from else "all"
    date_to_str   = str(date_to)   if date_to   else "all"
    filename = f"equipment_costs_{date_from_str}_{date_to_str}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export")
def export_costs(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    branch_id: Optional[int] = None,
    department_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _: dict = _auth,
):
    rows = _cost_rows(db, date_from, date_to, branch_id, department_id)
    rows = sorted(rows, key=lambda r: r.total_cost, reverse=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Затраты"

    headers = [
        "Инв.№", "Модель", "Филиал", "Отдел",
        "Затраты на ремонт", "Кол-во ремонтов",
        "Затраты на расходники", "Итого", "Цена отпечатка",
    ]
    ws.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r.inventory_number,
            r.model,
            r.branch or "",
            r.department or "",
            r.total_repair_cost,
            r.repair_count,
            r.total_consumable_cost,
            r.total_cost,
            r.cost_per_page if r.cost_per_page is not None else "",
        ])

    col_widths = [14, 24, 20, 20, 18, 16, 22, 14, 16]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    parts = ["costs"]
    if date_from:
        parts.append(str(date_from))
    if date_to:
        parts.append(str(date_to))
    filename = "_".join(parts) + ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/summary", response_model=SummaryResponse)
def summary(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    _: dict = _auth,
):
    total_active = (
        db.query(func.count(Device.id))
        .filter(Device.status == DeviceStatus.active)
        .scalar()
    )

    all_rows = _cost_rows(db, date_from, date_to, None, None)
    total_cost_period = round(sum(r.total_cost for r in all_rows), 2)
    top5 = sorted(all_rows, key=lambda r: r.total_cost, reverse=True)[:5]

    # Количество ремонтов по типам за период
    rbt_q = (
        db.query(
            RepairRecord.repair_type,
            func.count(RepairRecord.id).label("cnt"),
        )
        .group_by(RepairRecord.repair_type)
    )
    if date_from:
        rbt_q = rbt_q.filter(RepairRecord.date >= date_from)
    if date_to:
        rbt_q = rbt_q.filter(RepairRecord.date <= date_to)
    rbt = {row.repair_type: row.cnt for row in rbt_q.all()}

    return SummaryResponse(
        total_devices_active=total_active,
        total_cost_period=total_cost_period,
        top5_expensive=top5,
        repairs_by_type=RepairsByType(
            planned=rbt.get(RepairType.planned, 0),
            unplanned=rbt.get(RepairType.unplanned, 0),
            warranty=rbt.get(RepairType.warranty, 0),
        ),
    )
