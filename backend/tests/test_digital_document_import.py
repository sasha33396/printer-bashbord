import unittest
import re
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.exc import IntegrityError
from unittest.mock import patch

from digital_document_import import FIELDS, MAX_FILE_BYTES
from models import ElectronicSignature, MachineReadablePowerOfAttorney
import test_digital_documents as fixtures

SIGNATURE, POWER = fixtures.SIGNATURE, fixtures.POWER


class DigitalDocumentImportTests(unittest.TestCase):
    setUp = fixtures.DigitalDocumentsApiTests.setUp
    endpoint = fixtures.DigitalDocumentsApiTests.endpoint
    create = fixtures.DigitalDocumentsApiTests.create

    def workbook(self, signatures=None, powers=None, change=None):
        book = Workbook()
        book.remove(book.active)
        for kind, name, data in [('ecp', 'ЭЦП', signatures if signatures is not None else [SIGNATURE]),
                                  ('mchd', 'МЧД', powers if powers is not None else [POWER])]:
            sheet = book.create_sheet(name)
            sheet.append([label for _, label in FIELDS[kind]])
            for record in data:
                sheet.append([record.get(key) for key, _ in FIELDS[kind]])
        if change:
            change(book)
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        return buffer.getvalue()

    def rewrite_sheets(self, contents, rewrite):
        buffer = BytesIO()
        with ZipFile(BytesIO(contents)) as source, ZipFile(buffer, 'w', ZIP_DEFLATED) as target:
            for name in source.namelist():
                data = source.read(name)
                if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                    data = rewrite(data)
                target.writestr(name, data)
        return buffer.getvalue()

    def upload(self, contents, apply=False, filename='registers.xlsx'):
        return self.client.post('/api/digital-documents/import-xlsx', params={'apply': apply},
                                files={'file': (filename, contents, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})

    def assert_empty(self):
        self.assertEqual(self.client.get(self.endpoint('signatures')).json(), [])
        self.assertEqual(self.client.get(self.endpoint('powers-of-attorney')).json(), [])

    def test_preview_is_read_only_apply_imports_both_sheets_and_retry_skips_everything(self):
        file = self.workbook()
        preview = self.upload(file)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()['totals'], {'new': 2, 'created': 0, 'unchanged': 0, 'error': 0, 'conflict': 0})
        self.assertFalse(preview.json()['applied'])
        self.assert_empty()
        report = self.upload(file, apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['created'], 2)
        self.assertEqual(report['totals']['new'], 0)
        self.assertTrue(all(row['status'] == 'created' and row['existing_id'] for row in report['rows']))
        signature = self.client.get(self.endpoint('signatures')).json()[0]
        power = self.client.get(self.endpoint('powers-of-attorney')).json()[0]
        for key, value in SIGNATURE.items():
            self.assertEqual(signature[key], '2026-10-06T10:30:15Z' if key == 'revoked_at' else value)
        for key, value in POWER.items():
            self.assertEqual(power[key], value)
        retried = self.upload(file, apply=True).json()
        self.assertFalse(retried['applied'])
        self.assertEqual(retried['totals']['new'], 0)
        self.assertEqual(retried['totals']['unchanged'], 2)
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 1)

    def test_headers_any_order_title_blank_cells_and_optional_columns(self):
        def change(book):
            sheet = book['ЭЦП']
            book.remove(sheet)
            sheet = book.create_sheet(' эцп ')
            sheet.append(['Заголовок реестра'])
            sheet.append(['Лишний столбец', ' Действителен\nпо ', 'ФИО', 'Действителен c', 'Название\u00a0компании'])
            sheet.append(['ignore', '01.01.2027', 'Иванов Иван', '01.01.2026', 'ООО Компания'])
            sheet.append([None, None, None, None, None])
        report = self.upload(self.workbook(powers=[], change=change), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['created'], 1)
        self.assertEqual(report['rows'][0]['row'], 3)
        self.assertTrue(report['warnings'])
        record = self.client.get(self.endpoint('signatures')).json()[0]
        self.assertIsNone(record['inn'])
        self.assertEqual(record['valid_to'], '2027-01-01')

    def test_excel_dates_numeric_masks_and_local_revocation_time(self):
        def change(book):
            sheet = book['ЭЦП']
            sheet['B2'] = 12345678
            sheet['B2'].number_format = '0000000000'
            sheet['J2'] = datetime(2026, 1, 1)
            sheet['K2'] = 46388  # 2027-01-01 in Excel's default epoch.
            sheet['M2'] = datetime(2026, 10, 6, 14, 30, 15)
            sheet['O2'] = 123
            sheet['O2'].number_format = '00000000'
        report = self.upload(self.workbook(change=change), apply=True).json()
        self.assertTrue(report['applied'], report)
        record = self.client.get(self.endpoint('signatures')).json()[0]
        self.assertEqual(record['inn'], '0012345678')
        self.assertEqual(record['serial_number'], '00000123')
        self.assertEqual(record['valid_to'], '2027-01-01')
        self.assertEqual(record['revoked_at'], '2026-10-06T10:30:15Z')

    def test_error_in_second_sheet_does_not_block_correct_first_sheet(self):
        file = self.workbook(powers=[{**POWER, 'valid_to': '2025-01-01'}])
        for apply in [False, True]:
            report = self.upload(file, apply=apply).json()
            self.assertTrue(report['can_apply'])
            self.assertEqual(report['applied'], apply)
            self.assertEqual(report['totals']['error'], 1)
            issue = report['rows'][1]
            self.assertEqual((issue['sheet'], issue['row']), ('МЧД', 2))
            self.assertIn('Дата окончания', issue['message'])
            if apply:
                self.assertEqual(report['totals']['created'], 1)
                self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 1)
                self.assertEqual(self.client.get(self.endpoint('powers-of-attorney')).json(), [])
            else:
                self.assert_empty()

    def test_invalid_values_formulas_long_numeric_identifiers_and_blank_required_cells(self):
        for field, value in [('fingerprint', 12345678901234567890), ('inn', 123.5), ('email', 'broken'),
                             ('full_name', None), ('valid_from', '=TODAY()'), ('valid_to', 'not a date')]:
            with self.subTest(field=field):
                report = self.upload(self.workbook(signatures=[{**SIGNATURE, field: value}], powers=[]), apply=True).json()
                self.assertFalse(report['can_apply'], report)
                self.assertEqual(report['totals']['error'], 1)
                self.assertIn(dict(FIELDS['ecp'])[field], report['rows'][0]['message'])
                self.assert_empty()

    def test_duplicate_rows_in_file_are_added_once(self):
        report = self.upload(self.workbook(signatures=[SIGNATURE, SIGNATURE], powers=[POWER, POWER]), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['created'], 2)
        self.assertEqual(report['totals']['unchanged'], 2)
        self.assertIn('Повтор строки 2', report['rows'][1]['message'])

    def test_same_identity_different_data_is_skipped_while_other_sheet_imports(self):
        record = self.create('signatures', SIGNATURE)
        changed = {**SIGNATURE, 'position': 'Новая должность'}
        file = self.workbook(signatures=[changed])
        report = self.upload(file, apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['conflict'], 1)
        self.assertEqual(report['rows'][0]['existing_id'], record['id'])
        self.assertEqual(self.client.get(self.endpoint('signatures')).json(), [record])
        self.assertEqual(len(self.client.get(self.endpoint('powers-of-attorney')).json()), 1)

    def test_conflicting_duplicates_inside_file_are_both_withheld(self):
        report = self.upload(self.workbook(signatures=[SIGNATURE, {**SIGNATURE, 'full_name': 'Другой владелец'}]), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['conflict'], 2)
        self.assertEqual(report['totals']['created'], 1)
        self.assertEqual(self.client.get(self.endpoint('signatures')).json(), [])
        self.assertEqual(len(self.client.get(self.endpoint('powers-of-attorney')).json()), 1)

    def test_archived_records_are_skipped_without_restoring(self):
        record = self.create('signatures', SIGNATURE)
        self.client.delete(self.endpoint('signatures') + f"/{record['id']}")
        report = self.upload(self.workbook(), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['unchanged'], 1)
        self.assertIn('в архиве', report['rows'][0]['message'])
        self.assertEqual(self.client.get(self.endpoint('signatures')).json(), [])
        self.assertEqual(len(self.client.get(self.endpoint('signatures'), params={'archived': True}).json()), 1)

    def test_power_number_is_not_globally_unique(self):
        second = {**POWER, 'grantor_inn': '9999999999', 'grantor_name': 'Другой доверитель',
                  'fns_identifier': 'other-fns', 'edo_identifier': 'other-edo'}
        report = self.upload(self.workbook(signatures=[], powers=[POWER, second]), apply=True).json()
        self.assertTrue(report['applied'], report)
        self.assertEqual(report['totals']['created'], 2)

    def test_apply_rechecks_database_after_preview(self):
        file = self.workbook()
        self.assertEqual(self.upload(file).json()['totals']['new'], 2)
        self.create('powers-of-attorney', {**POWER, 'permissions': 'Изменено вручную'})
        applied = self.upload(file, apply=True).json()
        self.assertTrue(applied['applied'])
        self.assertEqual(applied['totals']['conflict'], 1)
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 1)
        self.assertEqual(applied['totals']['created'], 1)
        self.assertEqual(self.client.get(self.endpoint('powers-of-attorney')).json()[0]['permissions'], 'Изменено вручную')

    def test_insert_failure_rolls_back_first_sheet(self):
        original_flush = self.sessions.class_.flush
        def fail_on_power(db, *args, **kwargs):
            # Flush signature first, then fail second insert in the same transaction.
            signatures = [record for record in db.new if isinstance(record, ElectronicSignature)]
            if any(isinstance(record, MachineReadablePowerOfAttorney) for record in db.new):
                original_flush(db, objects=signatures)
                raise IntegrityError('forced', {}, Exception('forced'))
            return original_flush(db, *args, **kwargs)
        with patch.object(self.sessions.class_, 'flush', fail_on_power):
            with self.assertRaises(IntegrityError):
                self.upload(self.workbook(), apply=True)
        self.assert_empty()

    def test_invalid_file_missing_sheets_duplicate_headers_and_size_limit(self):
        cases = [(b'broken', 'register.xlsx'), (self.workbook(), 'register.xls'),
                 (b'x' * (MAX_FILE_BYTES + 1), 'register.xlsx'),
                 (self.workbook(change=lambda book: book.remove(book['МЧД'])), 'register.xlsx'),
                 (self.workbook(change=lambda book: setattr(book['ЭЦП']['B1'], 'value', 'Название компании')), 'register.xlsx'),
                 (self.workbook(signatures=[], powers=[]), 'register.xlsx')]
        for contents, filename in cases:
            with self.subTest(filename=filename, size=len(contents)):
                response = self.upload(contents, apply=True, filename=filename)
                self.assertEqual(response.status_code, 422, response.text)
                self.assertIsInstance(response.json()['detail'], str)
                self.assert_empty()

    def test_upload_requires_login(self):
        from routers.auth import get_current_user
        self.app.dependency_overrides.pop(get_current_user)
        self.assertEqual(self.upload(self.workbook(), apply=True).status_code, 401)

    def test_corrupted_sheet_xml_and_too_many_records_return_clear_errors(self):
        buffer = BytesIO()
        with ZipFile(BytesIO(self.workbook())) as source, ZipFile(buffer, 'w', ZIP_DEFLATED) as target:
            for name in source.namelist():
                contents = source.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    contents = contents.replace(b'</worksheet>', b'</invalid>')
                target.writestr(name, contents)
        response = self.upload(buffer.getvalue(), apply=True)
        self.assertEqual(response.status_code, 422, response.text)
        self.assertIsInstance(response.json()['detail'], str)
        self.assert_empty()
        with patch('digital_document_import.MAX_ROWS', 1):
            response = self.upload(self.workbook(signatures=[SIGNATURE, SIGNATURE]), apply=True)
        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn('максимум 1 строк', response.json()['detail'])
        self.assert_empty()

    def test_missing_underreported_and_inflated_dimensions_preserve_all_records(self):
        workbook = self.workbook()
        for dimension in [b'', b'<dimension ref="A1:A1"/>', b'<dimension ref="A1:XFD1048576"/>']:
            with self.subTest(dimension=dimension):
                contents = self.rewrite_sheets(workbook, lambda data: re.sub(rb'<dimension\b[^>]*/>', dimension, data))
                response = self.upload(contents)
                self.assertEqual(response.status_code, 200, response.text)
                report = response.json()
                self.assertTrue(report['can_apply'])
                self.assertEqual(report['totals']['new'], 2)
                self.assertEqual(report['rows'][0]['data']['serial_number'], '00000123')
                self.assertEqual(report['rows'][1]['data']['power_number'], '000045')
                self.assert_empty()

    def test_formatting_at_last_excel_cell_does_not_expand_table_or_block_apply(self):
        def change(book):
            for sheet in book:
                sheet['XFD1048576'].font = Font(bold=True)
                sheet['AAA1'].font = Font(italic=True)
        response = self.upload(self.workbook(change=change), apply=True)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()['applied'])
        self.assertEqual(response.json()['totals']['created'], 2)
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 1)
        self.assertEqual(len(self.client.get(self.endpoint('powers-of-attorney')).json()), 1)

    def test_wide_table_and_sparse_real_records_preserve_coordinates_and_validation(self):
        def change(book):
            sheet = book['ЭЦП']
            book.remove(sheet)
            sheet = book.create_sheet('ЭЦП')
            for column, (key, label) in enumerate(FIELDS['ecp'], 300):
                sheet.cell(2, column, label)
                sheet.cell(10021, column, SIGNATURE[key])
                sheet.cell(1048576, column, SIGNATURE[key] if key != 'full_name' else None)
        report = self.upload(self.workbook(change=change), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['created'], 2)
        self.assertEqual(report['totals']['error'], 1)
        self.assertEqual(report['rows'][0]['row'], 10021)
        self.assertEqual(report['rows'][1]['row'], 1048576)
        self.assertIn('ФИО', report['rows'][1]['message'])
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 1)
        self.assertEqual(len(self.client.get(self.endpoint('powers-of-attorney')).json()), 1)

    def test_conflicting_identifiers_pointing_to_multiple_records_are_not_merged(self):
        first = self.create('signatures', SIGNATURE)
        second = self.create('signatures', {**SIGNATURE, 'fingerprint': 'other', 'serial_number': 'other'})
        report = self.upload(self.workbook(signatures=[{**SIGNATURE, 'serial_number': 'other'}]), apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['conflict'], 1)
        self.assertEqual(self.client.get(self.endpoint('signatures')).json(), [first, second])

    def test_mixed_import_report_and_corrected_file_only_add_missing_records(self):
        existing = self.create('signatures', SIGNATURE)
        valid = {**SIGNATURE, 'fingerprint': 'valid-new', 'serial_number': '00000999', 'full_name': 'Новый сотрудник'}
        invalid = {**SIGNATURE, 'fingerprint': 'needs-fix', 'serial_number': '00000888', 'email': 'broken'}
        conflict = {**SIGNATURE, 'position': 'Расхождение'}
        power_invalid = {**POWER, 'fns_identifier': 'new-fns', 'edo_identifier': 'new-edo',
                         'power_number': '000055', 'valid_to': '2025-01-01'}
        file = self.workbook(signatures=[SIGNATURE, valid, invalid, conflict], powers=[POWER, power_invalid])
        preview = self.upload(file).json()
        self.assertTrue(preview['can_apply'])
        self.assertEqual(preview['totals']['new'], 2)
        report = self.upload(file, apply=True).json()
        self.assertTrue(report['applied'])
        self.assertEqual(report['totals']['created'], 2)
        self.assertEqual(report['totals']['error'], 2)
        self.assertEqual(self.client.get(self.endpoint('signatures') + f"/{existing['id']}").json(), existing)
        retried = self.upload(file, apply=True).json()
        self.assertFalse(retried['applied'])
        self.assertFalse(retried['can_apply'])
        self.assertEqual(retried['totals']['unchanged'], 2)
        corrected = self.workbook(signatures=[valid, {**invalid, 'email': 'fixed@example.com'}],
                                  powers=[POWER, {**power_invalid, 'valid_to': '2027-01-01'}])
        corrected_report = self.upload(corrected, apply=True).json()
        self.assertEqual(corrected_report['totals']['created'], 2)
        self.assertEqual(corrected_report['totals']['unchanged'], 2)
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 3)
        self.assertEqual(len(self.client.get(self.endpoint('powers-of-attorney')).json()), 2)
        response = self.client.post('/api/digital-documents/import-report-xlsx', json={**report, 'source_filename': 'Исходник.xlsx'})
        self.assertEqual(response.status_code, 200, response.text[:100])
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ['Итоги', 'ЭЦП', 'МЧД'])
        sheet = workbook['ЭЦП']
        headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(3, headers['Серийный номер']).value, '00000999')
        self.assertEqual(sheet.cell(4, headers['Email']).value, 'broken')
        self.assertEqual(sheet.cell(4, headers['Строка исходного Excel']).value, 4)
        self.assertIn('Email', sheet.cell(4, headers['Причина / пояснение']).value)
        self.assertEqual(sheet.cell(3, headers['Результат']).value, 'Добавлено')
        self.assertEqual(sheet.cell(5, headers['Результат']).value, 'Нужно сверить')
        self.assertEqual(workbook['Итоги']['B1'].value, 'Исходник.xlsx')
        self.assertEqual(len(self.client.get(self.endpoint('signatures')).json()), 3)

    def test_report_formulas_remain_text_and_export_requires_authentication(self):
        report = self.upload(self.workbook(signatures=[{**SIGNATURE, 'valid_from': '=TODAY()'}], powers=[])).json()
        response = self.client.post('/api/digital-documents/import-report-xlsx', json=report)
        self.assertEqual(response.status_code, 200, response.text[:100])
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook['ЭЦП']
        header_map = {cell.value: cell.column for cell in sheet[1]}
        cell = sheet.cell(2, header_map['Действителен с'])
        self.assertEqual(cell.value, '=TODAY()')
        self.assertEqual(cell.data_type, 's')
        self.assert_empty()
        from routers.auth import get_current_user
        self.app.dependency_overrides.pop(get_current_user)
        self.assertEqual(self.client.post('/api/digital-documents/import-report-xlsx', json=report).status_code, 401)

    def test_error_report_keeps_displayed_numeric_identifier_zeroes(self):
        def change(book):
            book['ЭЦП']['B2'] = 12345678
            book['ЭЦП']['B2'].number_format = '0000000000'
        report = self.upload(self.workbook(signatures=[{**SIGNATURE, 'email': 'broken'}], powers=[], change=change)).json()
        self.assertEqual(report['rows'][0]['source_data']['inn'], '0012345678')
        response = self.client.post('/api/digital-documents/import-report-xlsx', json=report)
        self.assertEqual(response.status_code, 200, response.text[:100])
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        sheet = workbook['ЭЦП']
        columns = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2, columns['ИНН']).value, '0012345678')
