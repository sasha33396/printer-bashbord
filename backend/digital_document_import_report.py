"""Downloadable import results with source values for manual correction."""
from collections import Counter
from io import BytesIO
from typing import Annotated, Literal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, StringConstraints

from digital_document_import import FIELDS, REGISTERS


Text = Annotated[str, StringConstraints(max_length=10000)]
SourceText = Annotated[str, StringConstraints(max_length=32767)]
STATUS_LABELS = {'new': 'Будет добавлено', 'created': 'Добавлено', 'unchanged': 'Уже есть / повтор',
                 'error': 'Ошибка', 'conflict': 'Нужно сверить'}


class ReportRow(BaseModel):
    kind: Literal['ecp', 'mchd']
    row: int = Field(ge=1)
    status: Literal['new', 'created', 'unchanged', 'error', 'conflict']
    message: Text = ''
    existing_id: int | None = Field(default=None, ge=1)
    data: dict[str, Text | None] | None = Field(default=None, max_length=15)
    source_data: dict[str, SourceText | None] = Field(default_factory=dict, max_length=15)


class ImportReportExport(BaseModel):
    mode: Literal['preview', 'apply']
    source_filename: Annotated[str, StringConstraints(max_length=500)] = ''
    rows: list[ReportRow] = Field(max_length=20000)
    warnings: list[Text] = Field(default_factory=list, max_length=20002)


def append_values(sheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(cell.value, str):
            # Formula-like source strings stay literal text. Identifiers retain zeroes.
            cell.data_type = 's'
            cell.number_format = '@'
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cells.append(cell)
    sheet.append(cells)


def report_workbook(report):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Итоги'
    append_values(summary, ['Отчёт импорта ЭЦП и МЧД', report.source_filename])
    append_values(summary, ['Режим', 'Результат импорта' if report.mode == 'apply' else 'Предварительная проверка'])
    totals = Counter(row.status for row in report.rows)
    for status, label in STATUS_LABELS.items():
        append_values(summary, [label, totals[status]])
    for warning in report.warnings:
        append_values(summary, ['Предупреждение', warning])
    summary.column_dimensions['A'].width = 32
    summary.column_dimensions['B'].width = 100
    for kind, (name, _, _) in REGISTERS.items():
        sheet = workbook.create_sheet(name)
        append_values(sheet, ['Строка исходного Excel', 'Результат', 'Причина / пояснение', 'ID записи в приложении',
                              *[label for _, label in FIELDS[kind]]])
        for row in report.rows:
            if row.kind != kind:
                continue
            values = row.source_data if row.status == 'error' else row.data or row.source_data
            append_values(sheet, [row.row, STATUS_LABELS[row.status], row.message, row.existing_id,
                                  *[values.get(key) for key, _ in FIELDS[kind]]])
        sheet.freeze_panes = 'E2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1677FF')
            sheet.column_dimensions[cell.column_letter].width = 24
        sheet.column_dimensions['C'].width = 85
    contents = BytesIO()
    workbook.save(contents)
    workbook.close()
    return contents.getvalue()
